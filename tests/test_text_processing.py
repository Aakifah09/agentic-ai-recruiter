from datetime import date

import pytest

from recruiter.documents import DocumentError, clean_text, extract_text
from recruiter.privacy import find_protected_terms, guess_name, jd_bias_flags, redact
from recruiter.skills import find_skills, normalise_skill, related_skills
from recruiter.textutils import (
    heading_section,
    is_grounded,
    merged_months,
    parse_point,
    split_sections,
    stated_years,
)


# ----------------------------------------------------------------------------- skills
@pytest.mark.parametrize("raw,expected", [
    ("python3", "Python"), ("Sklearn", "scikit-learn"), ("k8s", "Kubernetes"), ("Large Language Models", "LLMs"),
    ("React 18", "React"), ("GitHub Actions", "CI/CD"), ("Some Niche Tool", "Some Niche Tool"),
])
def test_normalise_skill(raw, expected):
    assert normalise_skill(raw) == expected


def test_find_skills_handles_symbols_and_boundaries():
    hits = find_skills("Built services in C++, C# and Node.js; Python/SQL pipelines; CI/CD with GitHub Actions.")
    assert {"C++", "C#", "Node.js", "Python", "SQL", "CI/CD", "Git"} <= set(hits)


def test_find_skills_avoids_common_word_false_positives():
    hits = find_skills("I like to go running and rest on weekends. R&D budget review.")
    assert not {"Go", "REST APIs", "R"} & set(hits)


def test_related_skills_are_symmetric():
    assert "Azure" in related_skills("AWS") and "AWS" in related_skills("Azure")


# ----------------------------------------------------------------------------- documents
def test_extract_pdf_and_docx_roundtrip(tmp_path):
    import docx
    import fitz

    text = "Jane Doe\nExperience\nData Analyst | Acme | Jan 2020 - Present\n- Built Power BI dashboards with SQL"
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((72, 72), text, fontsize=10)
    pdf_bytes = pdf.tobytes()
    assert "Power BI" in extract_text(pdf_bytes, "cv.pdf")

    d = docx.Document()
    for line in text.splitlines():
        d.add_paragraph(line)
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Skills", "Tableau, Python"
    path = tmp_path / "cv.docx"
    d.save(path)
    out = extract_text(path.read_bytes(), "cv.docx")
    assert "Power BI" in out and "Tableau, Python" in out


@pytest.mark.parametrize("data,name", [(b"hello", "cv.exe"), (b"", "cv.txt"), (b"short", "cv.txt"), (b"%PDF-broken", "cv.pdf")])
def test_extract_text_rejects_bad_input(data, name):
    with pytest.raises(DocumentError):
        extract_text(data, name)


def test_clean_text_normalises_bullets_and_whitespace():
    assert clean_text("• one\n\n\n\n▪  two\r\n") == "- one\n\n- two"


# ----------------------------------------------------------------------------- privacy
def test_redact_removes_pii_and_protected_lines_but_keeps_dates():
    raw = ("Maria Lopez\nmaria@example.com | +44 7700 900123 | linkedin.com/in/maria\nDate of Birth: 01/02/1990\n"
           "Nationality: Spanish\nGender: Female\nExperience\nAnalyst | Acme | 2019-2021\nShe built SQL reports. Maria led the team.")
    out = redact(raw)
    for leaked in ("maria@example.com", "7700", "linkedin", "1990", "Spanish", "Female", "Maria", "Lopez", "She "):
        assert leaked not in out.text
    assert "2019-2021" in out.text
    assert {"email", "phone", "url", "name", "protected characteristic", "pronoun"} <= set(out.redactions)


def test_guess_name_skips_headings():
    assert guess_name("Curriculum Vitae\nJohn Smith") is None
    assert guess_name("Name: Ana Silva\nSkills") == "Ana Silva"
    assert guess_name("Priya Natarajan | Data Scientist\n") == "Priya Natarajan"


def test_protected_terms_and_jd_bias():
    assert "family" in find_protected_terms("Do you have children?")
    assert find_protected_terms("Explain a race condition you debugged in retrieval augmented generation") == []
    flags = jd_bias_flags("We want a young rockstar who is a native English speaker.")
    assert len(flags) == 3


# ----------------------------------------------------------------------------- text utils
def test_date_parsing_and_overlap_merge():
    today = date(2025, 1, 1)
    assert parse_point("Present", today) == today
    assert parse_point("Sept 2021") == date(2021, 9, 1)
    assert parse_point("03/2020") == date(2020, 3, 1)
    intervals = [(date(2020, 1, 1), date(2021, 1, 1)), (date(2020, 6, 1), date(2022, 1, 1)), (date(2023, 1, 1), date(2023, 7, 1))]
    assert merged_months(intervals) == 24 + 6


def test_sections_and_headings():
    assert heading_section("Nice to have:") == "preferred"
    assert heading_section("WORK EXPERIENCE") == "experience"
    assert heading_section("Built a data warehouse with dbt") is None
    secs = split_sections("Intro\nSkills: Python, SQL\nEducation\nBSc")
    assert secs["skills"] == ["Python, SQL"] and secs["education"] == ["BSc"]


def test_stated_years():
    assert stated_years("3+ years of experience with Python") == [3.0]
    assert stated_years("Founded 25 years ago") == []


def test_grounding_rejects_hallucinated_quotes():
    src = "- Built RAG pipelines with LangChain and pgvector over 2M documents"
    assert is_grounded("Built RAG pipelines with LangChain and pgvector", src)
    assert is_grounded("built rag pipelines, with langchain and pgvector", src)
    assert not is_grounded("Led a team of 12 engineers building Kubernetes operators", src)
