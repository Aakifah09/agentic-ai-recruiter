"""Render a few sample resumes as PDF and DOCX to exercise the document parsers.

Usage:  python scripts/make_sample_docs.py
Writes to data/samples/uploads/ (not loaded automatically, use them in the upload widget).
"""

from __future__ import annotations

from pathlib import Path

import docx
import fitz

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "data" / "samples" / "resumes"
OUT = ROOT / "data" / "samples" / "uploads"


def to_pdf(text: str, path: Path) -> None:
    doc = fitz.open()
    page = doc.new_page()
    rect = fitz.Rect(50, 50, 545, 800)
    page.insert_textbox(rect, text, fontsize=9.5, fontname="helv")
    doc.save(path)


def to_docx(text: str, path: Path) -> None:
    d = docx.Document()
    lines = text.splitlines()
    d.add_heading(lines[0], 0)
    for line in lines[1:]:
        if not line.strip():
            continue
        if line.startswith("- "):
            d.add_paragraph(line[2:], style="List Bullet")
        elif len(line.split()) <= 3 and line[0].isupper() and "|" not in line:
            d.add_heading(line, 1)
        else:
            d.add_paragraph(line)
    d.save(path)


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    for name in ("r01_senior_ai_engineer", "r04_senior_data_analyst"):
        to_pdf((SRC / f"{name}.txt").read_text(encoding="utf-8"), OUT / f"{name}.pdf")
    for name in ("r05_qa_engineer", "r08_graduate_ai"):
        to_docx((SRC / f"{name}.txt").read_text(encoding="utf-8"), OUT / f"{name}.docx")
    print(f"Wrote {len(list(OUT.iterdir()))} files to {OUT}")


if __name__ == "__main__":
    main()
