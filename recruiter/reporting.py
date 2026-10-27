"""Exportable recruitment reports: shortlist table, Markdown, JSON and DOCX."""

from __future__ import annotations

import io
import json

import pandas as pd

from recruiter.schemas import EvidenceStrength, HumanReview, ShortlistReport

DECISION_LABELS = {"pending": "Pending review", "advance": "Advance to interview", "hold": "On hold",
                   "needs_info": "Needs more information"}


def shortlist_frame(report: ShortlistReport, reviews: dict[str, HumanReview] | None = None) -> pd.DataFrame:
    reviews = reviews or {}
    by_id = {c.candidate_id: c for c in report.candidates}
    rows = []
    for a in report.assessments:
        c = by_id[a.candidate_id]
        matched = [m.skill for m in a.requirement_matches if m.strength == EvidenceStrength.DEMONSTRATED]
        missing = [m.skill for m in a.requirement_matches if m.strength == EvidenceStrength.MISSING and m.importance.value == "essential"]
        review = reviews.get(a.candidate_id)
        rows.append({
            "rank": a.rank,
            "candidate": a.display_name,
            "candidate_id": a.candidate_id,
            "source_file": c.source_file,
            "overall_score": a.overall_score,
            "skill_score": a.skill_score,
            "essential_coverage": a.essential_coverage,
            "experience_score": a.experience_score,
            "qualification_score": a.qualification_score,
            "years_experience": c.total_years_experience,
            "confidence": a.confidence,
            "ai_recommendation": a.recommendation.value,
            "demonstrated_skills": ", ".join(matched),
            "missing_essentials": ", ".join(missing),
            "recruiter_decision": DECISION_LABELS[review.decision] if review else DECISION_LABELS["pending"],
            "recruiter_notes": review.notes if review else "",
        })
    return pd.DataFrame(rows)


def to_markdown(report: ShortlistReport, reviews: dict[str, HumanReview] | None = None) -> str:
    reviews = reviews or {}
    job = report.job
    by_id = {c.candidate_id: c for c in report.candidates}
    out = [
        f"# Recruitment shortlist - {job.title}",
        "",
        f"Run `{report.run_id}` · {report.created_at} · model: `{report.provider}`",
        "",
        f"> {report.disclaimer}",
        "",
        "## Job requirements",
        "",
        f"- Seniority: {job.seniority}; minimum experience: {job.min_years_experience or 'not stated'} years",
        f"- Education: {', '.join(job.education_requirements) or 'not stated'}",
        "",
        "| Skill | Importance | Weight |",
        "|---|---|---|",
        *[f"| {s.skill} | {s.importance.value} | {s.weight:.1%} |" for s in job.skills],
        "",
        "## Shortlist",
        "",
        "| Rank | Candidate | Score | Essential coverage | Experience (yrs) | Confidence | AI suggestion | Recruiter decision |",
        "|---|---|---|---|---|---|---|---|",
    ]
    for a in report.assessments:
        r = reviews.get(a.candidate_id)
        out.append(
            f"| {a.rank} | {a.display_name} | {a.overall_score:.0f} | {a.essential_coverage:.0f}% | "
            f"{by_id[a.candidate_id].total_years_experience} | {a.confidence} | {a.recommendation.value} | "
            f"{DECISION_LABELS[r.decision] if r else 'Pending review'} |"
        )
    for a in report.assessments:
        c = by_id[a.candidate_id]
        out += ["", f"## {a.rank}. {a.display_name}", "", a.rationale, ""]
        if a.strengths:
            out += ["**Strengths**", "", *[f"- {s}" for s in a.strengths], ""]
        if a.gaps:
            out += ["**Gaps / to explore**", "", *[f"- {g}" for g in a.gaps], ""]
        out += ["**Evidence by requirement**", "", "| Requirement | Evidence | Credit | Excerpt |", "|---|---|---|---|"]
        for m in a.requirement_matches:
            excerpt = (m.evidence[0][:140] if m.evidence else "-").replace("|", "/")
            out.append(f"| {m.skill} ({m.importance.value}) | {m.strength.value} | {m.credit:.0%} | {excerpt} |")
        plan = report.interview_plans.get(a.candidate_id)
        if plan and plan.questions:
            out += ["", "**Interview plan**", ""]
            out += [f"{i}. *[{q.purpose.replace('_', ' ')}]* {q.question}  \n   _Good answer:_ {q.what_good_looks_like}"
                    for i, q in enumerate(plan.questions, 1)]
        r = reviews.get(a.candidate_id)
        if r and r.notes:
            out += ["", f"**Recruiter notes ({r.reviewer or 'recruiter'}):** {r.notes}"]
        if c.redactions:
            out += ["", f"_Redacted before analysis: {', '.join(c.redactions)}._"]
    if report.compliance.issues:
        out += ["", "## Compliance review", ""]
        out += [f"- **{i.severity}** - {i.location}: {i.message}" for i in report.compliance.issues]
    return "\n".join(out) + "\n"


def to_json(report: ShortlistReport, reviews: dict[str, HumanReview] | None = None) -> str:
    data = report.model_dump(mode="json")
    data["human_reviews"] = {k: v.model_dump(mode="json") for k, v in (reviews or {}).items()}
    return json.dumps(data, indent=2)


def to_docx(report: ShortlistReport, reviews: dict[str, HumanReview] | None = None) -> bytes:
    import docx
    from docx.shared import Pt

    reviews = reviews or {}
    d = docx.Document()
    d.styles["Normal"].font.size = Pt(10)
    d.add_heading(f"Recruitment shortlist - {report.job.title}", 0)
    d.add_paragraph(f"Run {report.run_id} · {report.created_at} · model: {report.provider}")
    d.add_paragraph().add_run(report.disclaimer).italic = True

    d.add_heading("Shortlist", 1)
    df = shortlist_frame(report, reviews)[["rank", "candidate", "overall_score", "essential_coverage",
                                           "years_experience", "confidence", "ai_recommendation", "recruiter_decision"]]
    table = d.add_table(rows=1, cols=len(df.columns))
    table.style = "Light Grid Accent 1"
    for i, col in enumerate(df.columns):
        table.rows[0].cells[i].text = col.replace("_", " ").title()
    for _, row in df.iterrows():
        cells = table.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = f"{v:.0f}" if isinstance(v, float) and v > 1 else str(v)

    for a in report.assessments:
        d.add_heading(f"{a.rank}. {a.display_name} - {a.overall_score:.0f}/100", 1)
        d.add_paragraph(a.rationale)
        for label, items in (("Strengths", a.strengths), ("Gaps / to explore", a.gaps)):
            if items:
                d.add_paragraph(label, style="Heading 3")
                for s in items:
                    d.add_paragraph(s, style="List Bullet")
        plan = report.interview_plans.get(a.candidate_id)
        if plan and plan.questions:
            d.add_paragraph("Interview plan", style="Heading 3")
            for q in plan.questions:
                d.add_paragraph(f"[{q.purpose.replace('_', ' ')}] {q.question}", style="List Number")
        r = reviews.get(a.candidate_id)
        if r:
            d.add_paragraph(f"Recruiter decision: {DECISION_LABELS[r.decision]}. {r.notes}")
    if report.compliance.issues:
        d.add_heading("Compliance review", 1)
        for i in report.compliance.issues:
            d.add_paragraph(f"{i.severity.upper()} - {i.location}: {i.message}", style="List Bullet")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
